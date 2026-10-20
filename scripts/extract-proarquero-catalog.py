import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def slug(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-")


def text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def number(value: object) -> int | None:
    return int(value) if isinstance(value, (int, float)) and value > 0 else None


def category_for(section: str, fallback: str) -> str:
    value = section.upper()
    if "GUANTE" in value:
        return "Guantes"
    if "INDUMENTARIA" in value or "CALZA" in value:
        return "Indumentaria"
    if "PELOTA" in value:
        return "Pelotas"
    if "CALZADO" in value:
        return "Calzado"
    if "CANILLERA" in value or "ACCESOR" in value or "MEDIA" in value:
        return "Accesorios"
    return fallback


def brand_for(section: str) -> str:
    value = section.upper().strip()
    replacements = ["GUANTES", "INDUMENTARIA", "JUNIOR", "ADULTO", "LINEA", "PRODUCTOS", "CALZADO", "PELOTAS", "ACCESORIOS"]
    for word in replacements:
        value = value.replace(word, " ")
    value = re.sub(r"\s+", " ", value).strip()
    normalized = value.title() if value else "Proarquero"
    aliases = {"Vgfc": "VGFC", "Drb": "DRB", "Ch1": "CH1", "N1": "N1", "Ho Soccer": "HO Soccer", "Glo-Wpro": "Glow Pro", "Idumentari Glow Pro": "Glow Pro"}
    return aliases.get(normalized, normalized)


def variants_from(cells, sizes):
    variants = []
    for cell, size in zip(cells, sizes):
        stock = number(cell.value) or 0
        if stock or size:
            variants.append({"size": str(size), "stock": stock})
    return variants or [{"size": "Único", "stock": 0}]


def main(source: Path, destination: Path) -> None:
    sheet = load_workbook(source, data_only=True)["Hoja1"]
    products = []
    section = ""
    category = "Guantes"
    audience = "Adulto"
    sizes = [11, 10, 9, 8, 7, 6, 5, 4]

    for row in sheet.iter_rows(min_row=2, max_col=16):
        name = text(row[0].value)
        cash = number(row[4].value)
        if not name:
            continue
        if cash:
            product_slug = slug(f"{brand_for(section)}-{name}")
            product_variants = variants_from(row[7:15], sizes)
            products.append({
                "id": product_slug,
                "sourceRow": row[0].row,
                "name": name.title(),
                "slug": product_slug,
                "brand": brand_for(section),
                "section": section,
                "category": category,
                "audience": audience,
                "price": cash,
                "price3Installments": number(row[2].value),
                "price6Installments": number(row[3].value),
                "description": f"{brand_for(section)} · {category}. Consultá disponibilidad de talle y color.",
                "image": None,
                "isReferenceImage": False,
                "stock": sum(variant["stock"] for variant in product_variants),
                "color": None,
                "variantLabel": "Talle",
                "variantNote": "Stock actualizado desde la lista de Proarquero. Confirmá disponibilidad antes de finalizar tu pedido.",
                "variants": [{"id": f"{product_slug}-{variant['size']}", "label": f"Talle {variant['size']}", "size": variant["size"], "sourceStock": variant["stock"], "sourceAvailable": variant["stock"] > 0} for variant in product_variants],
            })
            continue
        section = name
        category = category_for(section, category)
        audience = "Junior" if "JUNIOR" in section.upper() else "Adulto"
        header_sizes = [number(cell.value) for cell in row[7:15]]
        if any(header_sizes):
            sizes = [size or fallback for size, fallback in zip(header_sizes, [11, 10, 9, 8, 7, 6, 5, 4])]

    apparel = [("Adulto", ["S", "M", "L", "XL"], 147, 170), ("Junior", ["8", "10", "12", "14", "16"], 177, 192)]
    for audience, sizes, start, end in apparel:
        for index in range(start, end + 1):
            name = text(sheet.cell(index, 7).value)
            if not name or "INDUMENTARIA" in name.upper():
                continue
            variants = []
            for offset, size in enumerate(sizes, start=8):
                variants.append({"size": size, "stock": number(sheet.cell(index, offset).value) or 0})
            product_slug = slug(f"vgfc-{name}-{audience}")
            products.append({
                "id": product_slug,
                "sourceRow": index,
                "name": name.title(),
                "slug": product_slug,
                "brand": "VGFC",
                "section": f"INDUMENTARIA {audience.upper()}",
                "category": "Indumentaria",
                "audience": audience,
                "price": 0,
                "price3Installments": None,
                "price6Installments": None,
                "description": f"VGFC · Indumentaria {audience.lower()}. Precio a confirmar.",
                "image": None,
                "isReferenceImage": False,
                "stock": sum(variant["stock"] for variant in variants),
                "color": None,
                "variantLabel": "Talle",
                "variantNote": "Precio a confirmar. Consultá disponibilidad antes de finalizar tu pedido.",
                "variants": [{"id": f"{product_slug}-{variant['size']}", "label": f"Talle {variant['size']}", "size": variant["size"], "sourceStock": variant["stock"], "sourceAvailable": variant["stock"] > 0} for variant in variants],
            })

    unique = {}
    for product in products:
        unique[product["slug"]] = product
    destination.write_text(json.dumps(list(unique.values()), ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"products": len(unique), "categories": sorted({product['category'] for product in unique.values()}), "brands": sorted({product['brand'] for product in unique.values()})}, ensure_ascii=False))


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
